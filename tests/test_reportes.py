from datetime import date, datetime, time
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

import Model as models
import routers.reportes as rutas
from reportes import datos_del_reporte, generar_excel, generar_pdf

DESDE, HASTA = date(2026, 10, 1), date(2026, 10, 31)


def fichaje(db, usuario, dia, hora, minuto=0, tipo="entrada", estatus="A tiempo"):
    # La hora se da en UTC: las 14:00 son las 08:00 en México.
    db.add(
        models.Asistencia(
            usuario_id=usuario.id,
            obra_id=usuario.obra_id,
            fecha_hora_fichaje=datetime(2026, 10, dia, hora, minuto),
            tipo=tipo,
            latitud_enviada=0,
            longitud_enviada=0,
            distancia_calculada_m=0,
            estatus_asistencia=estatus,
        )
    )
    db.commit()


def salida(db, usuario, dia, hora, minuto=0):
    fichaje(db, usuario, dia, hora, minuto, tipo="salida", estatus="Salida registrada")


def falta(db, usuario, dia):
    db.add(
        models.Falta(
            usuario_id=usuario.id, obra_id=usuario.obra_id, fecha=date(2026, 10, dia)
        )
    )
    db.commit()


def incidencia(db, usuario, dia, texto, prioridad="Media"):
    db.add(
        models.Incidencia(
            usuario_id=usuario.id,
            fecha=date(2026, 10, dia),
            motivo_texto=texto,
            categoria="Retardo",
            prioridad=prioridad,
        )
    )
    db.commit()


@pytest.fixture
def otra_obra(db):
    obra = models.Obra(
        nombre="Otra obra", latitud_centro=20, longitud_centro=-100, hora_entrada=time(8, 0)
    )
    db.add(obra)
    db.commit()
    return obra


@pytest.fixture
def ana(db, otra_obra):
    ana = models.Usuario(nombre="Ana", telefono="5215500000002", obra_id=otra_obra.id)
    db.add(ana)
    db.commit()
    return ana


@pytest.fixture
def periodo(db, usuario):
    # Día 5: jornada completa de 8 h 30 min. Día 6: retardo sin salida. Día 7: falta.
    fichaje(db, usuario, 5, 14)
    salida(db, usuario, 5, 22, 30)
    fichaje(db, usuario, 6, 15, estatus="Retardo")
    fichaje(db, usuario, 6, 16, estatus="Fuera de rango")
    falta(db, usuario, 7)
    incidencia(db, usuario, 6, "Voy a llegar tarde por el tráfico")
    incidencia(db, usuario, 7, "Me accidenté en la obra", prioridad="Alta")
    return datos_del_reporte(db, DESDE, HASTA)


def fila_de(datos, nombre):
    return next(f for f in datos["resumen"] if f["trabajador"] == nombre)


# --- Cálculo ---


def test_resumen_por_trabajador(periodo):
    assert fila_de(periodo, "Trabajador") == {
        "usuario_id": 1,
        "trabajador": "Trabajador",
        "obra": "Obra de prueba",
        "asistencias": 2,
        "retardos": 1,
        "faltas": 1,
        "horas": 8.5,
        "sin_salida": 1,
        "incidencias": 2,
    }
    assert periodo["totales"]["por_prioridad"] == {"Alta": 1, "Media": 1, "Baja": 0}
    assert periodo["totales"]["trabajadores"] == 1


def test_las_jornadas_llevan_entrada_salida_y_horas(periodo):
    completa, sin_salida = periodo["jornadas"]
    assert (completa["fecha"], completa["horas"]) == (date(2026, 10, 5), 8.5)
    assert completa["entrada"].strftime("%H:%M") == "08:00"
    assert completa["salida"].strftime("%H:%M") == "16:30"
    assert (sin_salida["estatus"], sin_salida["horas"]) == ("Retardo", None)


def test_los_intentos_fuera_de_rango_no_cuentan(db, usuario):
    fichaje(db, usuario, 5, 14, estatus="Fuera de rango")
    datos = datos_del_reporte(db, DESDE, HASTA)
    assert datos["jornadas"] == []
    assert fila_de(datos, "Trabajador")["asistencias"] == 0


def test_la_salida_de_madrugada_pertenece_a_su_dia_de_mexico(db, usuario):
    # Entrada a las 08:00 y salida a las 23:00 del día 5 (ya es día 6 en UTC).
    fichaje(db, usuario, 5, 14)
    salida(db, usuario, 6, 5)
    (jornada,) = datos_del_reporte(db, DESDE, HASTA)["jornadas"]
    assert (jornada["fecha"], jornada["horas"]) == (date(2026, 10, 5), 15.0)


def test_solo_entra_lo_del_periodo(db, usuario):
    fichaje(db, usuario, 5, 14)
    falta(db, usuario, 7)
    incidencia(db, usuario, 7, "aviso")
    datos = datos_del_reporte(db, date(2026, 10, 10), date(2026, 10, 20))
    assert datos["jornadas"] == datos["faltas"] == datos["incidencias"] == []
    # El personal asignado aparece aunque no tenga registros.
    assert fila_de(datos, "Trabajador")["asistencias"] == 0


def test_filtro_por_obra(db, usuario, ana, otra_obra):
    fichaje(db, usuario, 5, 14)
    fichaje(db, ana, 5, 14)
    incidencia(db, ana, 5, "aviso de Ana")

    datos = datos_del_reporte(db, DESDE, HASTA, otra_obra.id)
    assert [f["trabajador"] for f in datos["resumen"]] == ["Ana"]
    assert datos["obra"] == "Otra obra"
    assert len(datos["jornadas"]) == len(datos["incidencias"]) == 1
    assert len(datos_del_reporte(db, DESDE, HASTA)["resumen"]) == 2


def test_un_trabajador_dado_de_baja_conserva_su_historial(db, usuario):
    fichaje(db, usuario, 5, 14)
    usuario.activo = False
    db.commit()
    assert fila_de(datos_del_reporte(db, DESDE, HASTA), "Trabajador")["asistencias"] == 1
    assert datos_del_reporte(db, date(2026, 11, 1), date(2026, 11, 30))["resumen"] == []


def test_periodos_invalidos(db, obra):
    with pytest.raises(ValueError):
        datos_del_reporte(db, HASTA, DESDE)
    with pytest.raises(ValueError):
        datos_del_reporte(db, date(2025, 1, 1), date(2026, 10, 31))
    with pytest.raises(LookupError):
        datos_del_reporte(db, DESDE, HASTA, obra.id + 99)


# --- Archivos ---


def test_el_excel_trae_las_cuatro_hojas(periodo):
    libro = load_workbook(BytesIO(generar_excel(periodo)))
    assert libro.sheetnames == ["Resumen", "Jornadas", "Faltas", "Incidencias"]

    resumen = list(libro["Resumen"].values)
    assert resumen[0][:6] == (
        "Trabajador", "Obra", "Asistencias", "Retardos", "Faltas", "Horas trabajadas",
    )  # fmt: skip
    assert resumen[1][:6] == ("Trabajador", "Obra de prueba", 2, 1, 1, 8.5)

    jornadas = list(libro["Jornadas"].values)
    assert jornadas[1][3:] == ("08:00", "16:30", 8.5, "A tiempo")
    assert jornadas[2][3:] == ("09:00", None, None, "Retardo")
    assert libro["Faltas"].max_row == 2
    assert libro["Incidencias"].max_row == 3


def test_un_mensaje_no_puede_colarse_como_formula(db, usuario):
    incidencia(db, usuario, 6, '=HYPERLINK("http://malo.example","clic")')
    libro = load_workbook(BytesIO(generar_excel(datos_del_reporte(db, DESDE, HASTA))))
    celda = libro["Incidencias"]["F2"]
    assert celda.data_type == "s"
    assert celda.value.startswith("=HYPERLINK")


def test_el_pdf_se_genera_con_y_sin_datos(db, obra, periodo):
    assert generar_pdf(periodo).startswith(b"%PDF")
    vacio = datos_del_reporte(db, date(2026, 11, 1), date(2026, 11, 30), obra.id)
    assert generar_pdf(vacio).startswith(b"%PDF")


def test_el_pdf_aguanta_textos_con_simbolos(db, usuario):
    usuario.nombre = "Pérez & Hijos <b>"
    incidencia(db, usuario, 6, "llego tarde <a href='x'> & ya </para>")
    assert generar_pdf(datos_del_reporte(db, DESDE, HASTA)).startswith(b"%PDF")


# --- Rutas ---


def test_las_rutas_entregan_los_archivos(db, periodo):
    resumen = rutas.resumen(DESDE, HASTA, None, db)
    assert set(resumen) == {"desde", "hasta", "obra", "resumen", "totales"}

    excel = rutas.excel(DESDE, HASTA, None, db)
    assert excel.body.startswith(b"PK")
    assert "reporte_asistencia_2026-10-01_a_2026-10-31.xlsx" in (
        excel.headers["content-disposition"]
    )
    assert rutas.pdf(DESDE, HASTA, None, db).media_type == "application/pdf"


def test_las_rutas_explican_los_errores(db, obra):
    with pytest.raises(HTTPException) as error:
        rutas.resumen(HASTA, DESDE, None, db)
    assert error.value.status_code == 400
    with pytest.raises(HTTPException) as error:
        rutas.pdf(DESDE, HASTA, obra.id + 99, db)
    assert error.value.status_code == 404
