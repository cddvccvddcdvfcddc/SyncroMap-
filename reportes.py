from datetime import date, timedelta
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

import Model as models
from utils.horario import MESES, a_hora_local, ahora_utc, inicio_del_dia_utc

EMPRESA = "MONTAJES IIMYRSA"
FUERA_DE_RANGO = "Fuera de rango"
PRIORIDADES = ("Alta", "Media", "Baja")

# Un reporte abarca como máximo un año; más que eso suele ser un error de captura.
MAXIMO_DE_DIAS = 366

# En la gráfica del PDF caben bien hasta este número de trabajadores.
TRABAJADORES_EN_LA_GRAFICA = 12


def dia_legible(dia: date) -> str:
    return f"{dia.day} de {MESES[dia.month - 1]} de {dia.year}"


def _horas(entrada, salida) -> Optional[float]:
    # Horas entre la entrada y la salida; vacío si falta alguna de las dos.
    if entrada is None or salida is None or salida <= entrada:
        return None
    return round((salida - entrada).total_seconds() / 3600, 2)


def datos_del_reporte(
    db: Session, desde: date, hasta: date, obra_id: Optional[int] = None
) -> dict:
    # Reúne todo lo del periodo: jornadas, faltas, incidencias y el resumen por trabajador.
    if hasta < desde:
        raise ValueError("La fecha final no puede ser anterior a la inicial.")
    if (hasta - desde).days >= MAXIMO_DE_DIAS:
        raise ValueError("El periodo no puede ser mayor a un año.")

    obras = {obra.id: obra for obra in db.query(models.Obra).all()}
    if obra_id is not None and obra_id not in obras:
        raise LookupError("Obra no encontrada")
    usuarios = {usuario.id: usuario for usuario in db.query(models.Usuario).all()}

    def de_la_obra(usuario_id: int) -> bool:
        return obra_id is None or usuarios[usuario_id].obra_id == obra_id

    # 1. Jornadas: una por trabajador y día, con su entrada y su salida válidas.
    fichajes = db.query(models.Asistencia).filter(
        models.Asistencia.estatus_asistencia != FUERA_DE_RANGO,
        models.Asistencia.fecha_hora_fichaje >= inicio_del_dia_utc(desde),
        models.Asistencia.fecha_hora_fichaje
        < inicio_del_dia_utc(hasta + timedelta(days=1)),
    )
    if obra_id is not None:
        fichajes = fichajes.filter(models.Asistencia.obra_id == obra_id)

    por_dia: dict = {}
    for fichaje in fichajes.order_by(models.Asistencia.fecha_hora_fichaje).all():
        momento = a_hora_local(fichaje.fecha_hora_fichaje)
        jornada = por_dia.setdefault(
            (fichaje.usuario_id, momento.date()),
            {"obra_id": fichaje.obra_id, "entrada": None, "salida": None, "estatus": None},
        )
        if fichaje.tipo == "entrada" and jornada["entrada"] is None:
            jornada["entrada"] = momento
            jornada["estatus"] = fichaje.estatus_asistencia
        elif fichaje.tipo == "salida":
            jornada["salida"] = momento

    jornadas = [
        {
            "usuario_id": usuario_id,
            "fecha": dia,
            "trabajador": usuarios[usuario_id].nombre,
            "obra": obras[jornada["obra_id"]].nombre,
            "entrada": jornada["entrada"],
            "salida": jornada["salida"],
            "horas": _horas(jornada["entrada"], jornada["salida"]),
            "estatus": jornada["estatus"],
        }
        for (usuario_id, dia), jornada in por_dia.items()
        if jornada["entrada"] is not None
    ]
    jornadas.sort(key=lambda j: (j["fecha"], j["trabajador"]))

    # 2. Faltas (las registra la tarea programada al terminar cada día).
    consulta = db.query(models.Falta).filter(
        models.Falta.fecha >= desde, models.Falta.fecha <= hasta
    )
    if obra_id is not None:
        consulta = consulta.filter(models.Falta.obra_id == obra_id)
    faltas = [
        {
            "usuario_id": falta.usuario_id,
            "fecha": falta.fecha,
            "trabajador": usuarios[falta.usuario_id].nombre,
            "obra": obras[falta.obra_id].nombre,
        }
        for falta in consulta.order_by(models.Falta.fecha).all()
    ]
    faltas.sort(key=lambda f: (f["fecha"], f["trabajador"]))

    # 3. Incidencias; no guardan obra, así que cuentan las del personal de la obra.
    incidencias = [
        {
            "usuario_id": incidencia.usuario_id,
            "fecha": incidencia.fecha,
            "trabajador": usuarios[incidencia.usuario_id].nombre,
            "categoria": incidencia.categoria,
            "prioridad": incidencia.prioridad,
            "estatus_revision": incidencia.estatus_revision,
            "motivo": incidencia.motivo_texto,
        }
        for incidencia in db.query(models.Incidencia)
        .filter(models.Incidencia.fecha >= desde, models.Incidencia.fecha <= hasta)
        .order_by(models.Incidencia.fecha, models.Incidencia.id)
        .all()
        if de_la_obra(incidencia.usuario_id)
    ]

    # 4. Resumen: el personal activo asignado, más quien tenga registros en el periodo.
    con_registros = {r["usuario_id"] for r in jornadas + faltas + incidencias}
    asignados = {
        usuario.id
        for usuario in usuarios.values()
        if usuario.activo and usuario.obra_id is not None and de_la_obra(usuario.id)
    }

    resumen = []
    for usuario_id in con_registros | asignados:
        usuario = usuarios[usuario_id]
        suyas = [j for j in jornadas if j["usuario_id"] == usuario_id]
        obra = obras.get(usuario.obra_id)
        resumen.append(
            {
                "usuario_id": usuario_id,
                "trabajador": usuario.nombre,
                "obra": obra.nombre if obra else "Sin obra",
                "asistencias": len(suyas),
                "retardos": sum(j["estatus"] == "Retardo" for j in suyas),
                "faltas": sum(f["usuario_id"] == usuario_id for f in faltas),
                "horas": round(sum(j["horas"] or 0 for j in suyas), 2),
                "sin_salida": sum(j["salida"] is None for j in suyas),
                "incidencias": sum(i["usuario_id"] == usuario_id for i in incidencias),
            }
        )
    resumen.sort(key=lambda fila: fila["trabajador"].lower())

    columnas = ("asistencias", "retardos", "faltas", "horas", "sin_salida", "incidencias")
    totales = {c: round(sum(fila[c] for fila in resumen), 2) for c in columnas}
    totales["trabajadores"] = len(resumen)
    totales["por_prioridad"] = {
        prioridad: sum(i["prioridad"] == prioridad for i in incidencias)
        for prioridad in PRIORIDADES
    }

    return {
        "desde": desde,
        "hasta": hasta,
        "obra": obras[obra_id].nombre if obra_id is not None else "Todas las obras",
        "generado": a_hora_local(ahora_utc()),
        "resumen": resumen,
        "totales": totales,
        "jornadas": jornadas,
        "faltas": faltas,
        "incidencias": incidencias,
    }


def nombre_de_archivo(datos: dict, extension: str) -> str:
    return f"reporte_asistencia_{datos['desde']}_a_{datos['hasta']}.{extension}"


# ==========================================
# EXCEL (para el área de nóminas)
# ==========================================

RELLENO_DE_ENCABEZADO = PatternFill("solid", fgColor="0F766E")


def _hoja(libro: Workbook, titulo: str, columnas: list, filas: list):
    # columnas: (título, ancho). Los textos nunca se interpretan como fórmulas.
    hoja = libro.create_sheet(titulo)
    hoja.append([nombre for nombre, _ in columnas])
    for indice, (_, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=indice)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = RELLENO_DE_ENCABEZADO
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[celda.column_letter].width = ancho

    for numero, fila in enumerate(filas, start=2):
        for indice, valor in enumerate(fila, start=1):
            celda = hoja.cell(row=numero, column=indice, value=valor)
            if isinstance(valor, str):
                # Un mensaje que empiece con "=" se guarda como texto, no como fórmula.
                celda.data_type = "s"
            elif isinstance(valor, date):
                celda.number_format = "DD/MM/YYYY"
    hoja.freeze_panes = "A2"
    return hoja


def _hora(momento) -> str:
    return momento.strftime("%H:%M") if momento else ""


def generar_excel(datos: dict) -> bytes:
    libro = Workbook()
    libro.remove(libro.active)
    totales = datos["totales"]

    _hoja(
        libro,
        "Resumen",
        [("Trabajador", 32), ("Obra", 26), ("Asistencias", 13), ("Retardos", 11),
         ("Faltas", 9), ("Horas trabajadas", 17), ("Días sin salida", 16),
         ("Incidencias", 13)],  # fmt: skip
        [
            [f["trabajador"], f["obra"], f["asistencias"], f["retardos"], f["faltas"],
             f["horas"], f["sin_salida"], f["incidencias"]]  # fmt: skip
            for f in datos["resumen"]
        ]
        + [
            [],
            ["TOTAL", "", totales["asistencias"], totales["retardos"], totales["faltas"],
             totales["horas"], totales["sin_salida"], totales["incidencias"]],  # fmt: skip
            [],
            ["Empresa", EMPRESA],
            ["Periodo", f"{dia_legible(datos['desde'])} al {dia_legible(datos['hasta'])}"],
            ["Obra", datos["obra"]],
            ["Generado", datos["generado"].strftime("%d/%m/%Y %H:%M")],
        ],
    )
    _hoja(
        libro,
        "Jornadas",
        [("Fecha", 12), ("Trabajador", 32), ("Obra", 26), ("Entrada", 10),
         ("Salida", 10), ("Horas", 9), ("Estatus", 12)],  # fmt: skip
        [
            [j["fecha"], j["trabajador"], j["obra"], _hora(j["entrada"]),
             _hora(j["salida"]), j["horas"], j["estatus"]]  # fmt: skip
            for j in datos["jornadas"]
        ],
    )
    _hoja(
        libro,
        "Faltas",
        [("Fecha", 12), ("Trabajador", 32), ("Obra", 26)],
        [[f["fecha"], f["trabajador"], f["obra"]] for f in datos["faltas"]],
    )
    _hoja(
        libro,
        "Incidencias",
        [("Fecha", 12), ("Trabajador", 32), ("Categoría", 13), ("Prioridad", 11),
         ("Revisión", 12), ("Mensaje", 80)],  # fmt: skip
        [
            [i["fecha"], i["trabajador"], i["categoria"], i["prioridad"] or "",
             i["estatus_revision"], i["motivo"]]  # fmt: skip
            for i in datos["incidencias"]
        ],
    )

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ==========================================
# PDF (reporte ejecutivo)
# ==========================================

VERDE = colors.HexColor("#0F766E")
SERIES = (
    ("Asistencias", "asistencias", colors.HexColor("#059669")),
    ("Retardos", "retardos", colors.HexColor("#D97706")),
    ("Faltas", "faltas", colors.HexColor("#DC2626")),
)


def _tabla(
    filas: list, anchos: list, alinear: str = "numeros", resaltar_ultima: bool = False
) -> Table:
    # alinear: "numeros" (cifras a la derecha desde la tercera columna), "CENTER" o "LEFT".
    alineacion = (
        ("ALIGN", (2, 0), (-1, -1), "RIGHT")
        if alinear == "numeros"
        else ("ALIGN", (0, 0), (-1, -1), alinear)
    )
    tabla = Table(filas, colWidths=anchos, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), VERDE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        alineacion,
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]
    if resaltar_ultima:
        estilo.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
    tabla.setStyle(TableStyle(estilo))
    return tabla


def _grafica(resumen: list) -> Drawing:
    # Barras por trabajador: asistencias, retardos y faltas.
    filas = sorted(resumen, key=lambda f: -(f["asistencias"] + f["faltas"]))
    filas = filas[:TRABAJADORES_EN_LA_GRAFICA]

    dibujo = Drawing(17 * cm, 7 * cm)
    grafica = VerticalBarChart()
    grafica.x, grafica.y = 1 * cm, 1.6 * cm
    grafica.width, grafica.height = 13 * cm, 4.8 * cm
    grafica.data = [[fila[campo] for fila in filas] for _, campo, _ in SERIES]
    grafica.categoryAxis.categoryNames = [fila["trabajador"][:16] for fila in filas]
    grafica.categoryAxis.labels.fontSize = 7
    grafica.categoryAxis.labels.fontName = "Helvetica"
    grafica.categoryAxis.labels.angle = 25
    grafica.categoryAxis.labels.boxAnchor = "ne"
    grafica.valueAxis.valueMin = 0
    grafica.valueAxis.labels.fontSize = 7
    grafica.valueAxis.labels.fontName = "Helvetica"
    maximo = max((max(serie) for serie in grafica.data), default=0)
    grafica.valueAxis.valueMax = max(maximo, 1)
    grafica.valueAxis.valueStep = max(1, round(maximo / 5))
    grafica.barSpacing = 1
    grafica.groupSpacing = 8
    for indice, (_, _, color) in enumerate(SERIES):
        grafica.bars[indice].fillColor = color
        grafica.bars[indice].strokeColor = None
    dibujo.add(grafica)

    leyenda = Legend()
    leyenda.x, leyenda.y = 14.6 * cm, 5.8 * cm
    leyenda.fontSize = 8
    leyenda.fontName = "Helvetica"
    leyenda.alignment = "right"
    leyenda.colorNamePairs = [(color, nombre) for nombre, _, color in SERIES]
    dibujo.add(leyenda)
    return dibujo


def generar_pdf(datos: dict) -> bytes:
    estilos = getSampleStyleSheet()
    normal, titulo, seccion = estilos["Normal"], estilos["Title"], estilos["Heading2"]
    seccion.textColor = VERDE
    seccion.fontSize = 12
    celda = ParagraphStyle("celda", parent=normal, fontSize=8.5, leading=10.5)
    totales = datos["totales"]

    def escapar(valor) -> str:
        # Los textos libres se escapan para que no rompan el formato del PDF.
        return str(valor).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def texto(valor) -> Paragraph:
        return Paragraph(escapar(valor), celda)

    contenido = [
        Paragraph("Reporte de asistencia", titulo),
        Paragraph(
            f"<b>{EMPRESA}</b> · SyncroMap<br/>"
            f"Periodo: {dia_legible(datos['desde'])} al {dia_legible(datos['hasta'])}<br/>"
            f"Obra: {escapar(datos['obra'])}<br/>"
            f"Generado el {datos['generado'].strftime('%d/%m/%Y a las %H:%M')}",
            normal,
        ),
        Spacer(1, 0.5 * cm),
        _tabla(
            [
                ["Trabajadores", "Asistencias", "Retardos", "Faltas", "Horas", "Incidencias"],
                [totales["trabajadores"], totales["asistencias"], totales["retardos"],
                 totales["faltas"], f"{totales['horas']:.2f}", totales["incidencias"]],  # fmt: skip
            ],
            [2.9 * cm] * 6,
            alinear="CENTER",
        ),
    ]

    if datos["resumen"]:
        contenido += [
            Paragraph("Asistencias, retardos y faltas por trabajador", seccion),
            _grafica(datos["resumen"]),
        ]
        if len(datos["resumen"]) > TRABAJADORES_EN_LA_GRAFICA:
            contenido.append(
                Paragraph(
                    f"La gráfica muestra a los {TRABAJADORES_EN_LA_GRAFICA} trabajadores "
                    "con más registros; la tabla los incluye a todos.",
                    normal,
                )
            )

    contenido += [
        Paragraph("Resumen por trabajador", seccion),
        _tabla(
            [["Trabajador", "Obra", "Asist.", "Retardos", "Faltas", "Horas", "Incid."]]
            + [
                [texto(f["trabajador"]), texto(f["obra"]), f["asistencias"], f["retardos"],
                 f["faltas"], f"{f['horas']:.2f}", f["incidencias"]]  # fmt: skip
                for f in datos["resumen"]
            ]
            + [
                ["Total", "", totales["asistencias"], totales["retardos"], totales["faltas"],
                 f"{totales['horas']:.2f}", totales["incidencias"]]  # fmt: skip
            ],
            [5.2 * cm, 4.3 * cm, 1.5 * cm, 1.8 * cm, 1.5 * cm, 1.7 * cm, 1.5 * cm],
            resaltar_ultima=True,
        ),
    ]

    contenido.append(Paragraph("Incidencias del periodo", seccion))
    if datos["incidencias"]:
        por_prioridad = totales["por_prioridad"]
        contenido += [
            Paragraph(
                " · ".join(f"{p}: <b>{por_prioridad[p]}</b>" for p in PRIORIDADES), normal
            ),
            Spacer(1, 0.3 * cm),
            _tabla(
                [["Fecha", "Trabajador", "Categoría", "Prioridad", "Mensaje"]]
                + [
                    [i["fecha"].strftime("%d/%m/%Y"), texto(i["trabajador"]), i["categoria"],
                     i["prioridad"] or "—", texto(i["motivo"])]  # fmt: skip
                    for i in datos["incidencias"]
                ],
                [2 * cm, 3.8 * cm, 2 * cm, 1.8 * cm, 7.9 * cm],
                alinear="LEFT",
            ),
        ]
    else:
        contenido.append(Paragraph("No se reportaron incidencias en el periodo.", normal))

    def pie(lienzo, documento):
        lienzo.setFont("Helvetica", 8)
        lienzo.setFillColor(colors.HexColor("#64748B"))
        lienzo.drawString(2 * cm, 1.2 * cm, f"{EMPRESA} · Reporte de asistencia")
        lienzo.drawRightString(letter[0] - 2 * cm, 1.2 * cm, f"Página {documento.page}")

    salida = BytesIO()
    SimpleDocTemplate(
        salida,
        pagesize=letter,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=1.8 * cm,
        bottomMargin=2 * cm,
        title="Reporte de asistencia",
        author=EMPRESA,
    ).build(contenido, onFirstPage=pie, onLaterPages=pie)
    return salida.getvalue()
